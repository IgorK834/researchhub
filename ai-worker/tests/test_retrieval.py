import json
from io import BytesIO
from pathlib import Path
from uuid import uuid4
import pytest
from docx import Document
from pydantic import ValidationError
from researchhub_worker.retrieval import chunk_extraction
from researchhub_worker.retrieval.contracts import ChunkingConfig, RetrievalChunkSet, RetrievalChunk, SourceSpan, digest
from researchhub_worker.contracts import SourceIngestResult
from researchhub_worker.parsing import SourceParser
from test_parsing import command, pdf_bytes

FIXTURES = Path(__file__).resolve().parents[2] / 'contracts'


def parse(kind, data):
    result = SourceParser(downloader=lambda *_: data)(command(kind))
    assert result.status == 'SUCCEEDED', result.failure
    return result


def docx(blocks):
    document = Document()
    for level, text in blocks:
        document.add_heading(text, level) if level else document.add_paragraph(text)
    stream = BytesIO(); document.save(stream)
    return parse('DOCX', stream.getvalue())


def test_short_sections_merge_body_with_heading_and_keep_titles():
    result = docx([(1, 'Methods'), (0, 'Short paragraph.'), (0, 'Another paragraph.'), (1, 'Results'), (0, 'Evidence.')])
    chunks = result.retrieval.chunks
    assert len(chunks) == 2
    assert chunks[0].content == 'Methods\n\nShort paragraph.\n\nAnother paragraph.'
    assert chunks[0].section_title == 'Methods' and chunks[1].section_title == 'Results'
    assert len(chunks[0].spans) == 3


def test_parent_heading_does_not_become_an_isolated_fragment():
    result = docx([(1, 'Methods'), (2, 'Measurements'), (0, 'Measured body'), (1, 'Appendix')])
    chunks = result.retrieval.chunks
    assert chunks[0].content == 'Methods\n\nMeasurements\n\nMeasured body'
    assert chunks[0].section_title == 'Measurements'
    # A final heading has no safely associated body; retain its provenance.
    assert chunks[1].content == 'Appendix'


def test_oversized_section_splits_at_paragraphs_with_overlap_and_context():
    result = docx([(1, 'Methods'), (0, 'First paragraph ' * 10), (0, 'Second paragraph ' * 10), (0, 'Third paragraph ' * 10)])
    config = ChunkingConfig(max_characters=160, overlap_characters=20, min_characters=40)
    chunks = chunk_extraction(result, config).chunks
    assert len(chunks) > 3
    assert all(len(chunk.content) <= 160 and chunk.section_title == 'Methods' for chunk in chunks)
    assert chunks[0].content.startswith('Methods\n\n')
    assert any(a.spans[-1].character_end > b.spans[0].character_start for a, b in zip(chunks, chunks[1:]))


def test_pdf_pages_never_merge_even_when_tiny_and_blank_pages_remain_traceable():
    result = parse('PDF', pdf_bytes(['Page A', 'Page B', '']))
    chunks = result.retrieval.chunks
    assert [(chunk.page_start, chunk.page_end) for chunk in chunks] == [(1, 1), (2, 2)]
    assert result.structure.pages[2].page_number == 3
    assert all(chunk.spans[0].unit_id == f'unit-{chunk.chunk_index}' for chunk in chunks)


def test_unicode_hard_windows_rebalance_tail_and_reconstruct_original_ranges():
    text = 'λ😀' * 37 + 'x'
    result = parse('TXT', text.encode())
    config = ChunkingConfig(max_characters=64, overlap_characters=8, min_characters=24)
    chunks = chunk_extraction(result, config).chunks
    assert [len(chunk.content) for chunk in chunks] == [59, 24]
    for chunk in chunks:
        span = chunk.spans[0]
        assert text[span.character_start:span.character_end] == chunk.content
        assert chunk.content_hash == digest(chunk.content)
    assert chunks[0].content[-8:] == chunks[1].content[:8]


def test_unicode_whitespace_trimming_preserves_all_non_whitespace_source_text():
    result = SourceIngestResult.model_validate(json.loads((FIXTURES/'retrieval/v1/unicode-whitespace-extraction.json').read_text()))
    assert result.retrieval == chunk_extraction(result, result.retrieval.config)
    text = result.chunks[0].text
    covered = set()
    for chunk in result.retrieval.chunks:
        span = chunk.spans[0]
        assert text[span.character_start:span.character_end] == chunk.content
        covered.update(range(span.character_start, span.character_end))
    assert all(index in covered or character.isspace() for index, character in enumerate(text))
    assert len(covered) < len(text), 'The fixture must exercise whitespace gaps between windows'


def test_determinism_stable_reprocessing_ids_and_incompatible_config_versions():
    result = parse('TXT', ('body word ' * 300).encode())
    first = chunk_extraction(result)
    assert first == chunk_extraction(result.model_copy(update={'job_id': uuid4()}))
    changed = chunk_extraction(result, ChunkingConfig(overlap_characters=0))
    assert first.processing_version != changed.processing_version
    assert set(c.chunk_id for c in first.chunks).isdisjoint(c.chunk_id for c in changed.chunks)
    workspace = chunk_extraction(result.model_copy(update={'workspace_id': uuid4()}))
    assert first.chunks[0].chunk_id != workspace.chunks[0].chunk_id


def test_shared_schema_round_trip_and_exact_search_projection():
    for filename in ['source-ingest-result-success.json', 'source-ingest-result-workbook.json']:
        result = SourceIngestResult.model_validate(json.loads((FIXTURES/'processing/v4'/filename).read_text()))
        assert result.retrieval == chunk_extraction(result)
    raw = json.loads((FIXTURES/'retrieval/v1/chunk-set.json').read_text())
    result = RetrievalChunkSet.model_validate(raw)
    assert result.model_dump(mode='json', by_alias=True) == raw
    assert result.chunks[0].search_document() == raw['chunks'][0]
    schema = RetrievalChunk.model_json_schema(by_alias=True)
    assert schema == json.loads((FIXTURES/'retrieval/v1/chunk.schema.json').read_text())
    search_schema = json.loads((FIXTURES/'retrieval/v1/search-document.schema.json').read_text())
    assert search_schema['properties'] == schema['properties'] and search_schema['required'] == schema['required']
    assert search_schema['x-researchhub-search']['mandatorySecurityFilters'] == ['workspaceId']


@pytest.mark.parametrize('values', [dict(max_characters=31), dict(overlap_characters=-1), dict(max_characters=64,overlap_characters=33,min_characters=1), dict(max_characters=64,overlap_characters=8,min_characters=57), dict(version='latest')])
def test_refuses_unsafe_chunk_configuration(values):
    with pytest.raises(ValidationError): ChunkingConfig(**values)


def test_runtime_configuration_is_pinned_and_validated(monkeypatch):
    monkeypatch.setenv('AI_WORKER_CHUNK_MAX_CHARACTERS', '80')
    monkeypatch.setenv('AI_WORKER_CHUNK_OVERLAP_CHARACTERS', '10')
    monkeypatch.setenv('AI_WORKER_CHUNK_MIN_CHARACTERS', '20')
    assert ChunkingConfig.from_env().max_characters == 80
    monkeypatch.setenv('AI_WORKER_CHUNK_OVERLAP_CHARACTERS', '41')
    with pytest.raises(ValidationError): SourceParser()


@pytest.mark.parametrize('mutation', [
    lambda raw: raw.update(processingVersion='old'),
    lambda raw: raw['chunks'][0].update(workspaceId=str(uuid4())),
    lambda raw: raw['chunks'][0].update(chunkId='0'*64),
    lambda raw: raw['chunks'][0].update(contentHash='0'*64),
    lambda raw: raw['chunks'][0].update(pageEnd=0),
    lambda raw: raw['chunks'][0].update(pageEnd=None),
    lambda raw: raw['chunks'][0].update(chunkIndex=2),
    lambda raw: raw['chunks'][0]['spans'][0].update(characterEnd=0),
])
def test_rejects_forged_identities_versions_hashes_and_locations(mutation):
    raw = json.loads((FIXTURES/'retrieval/v1/chunk-set.json').read_text()); mutation(raw)
    with pytest.raises(ValidationError): RetrievalChunkSet.model_validate(raw)


def test_a_long_heading_stays_with_body_when_the_window_can_fit_both():
    result = docx([(1, 'H' * 320), (0, 'body ' * 100)])
    chunks = chunk_extraction(result, ChunkingConfig(max_characters=512, overlap_characters=0, min_characters=64)).chunks
    assert 'body' in chunks[0].content and len(chunks[0].spans) == 2
    # If a heading itself exceeds the cap, its windows still carry the right title.
    short = chunk_extraction(result, ChunkingConfig(max_characters=128, overlap_characters=0, min_characters=32))
    assert all(chunk.section_title == 'H' * 320 for chunk in short.chunks)
