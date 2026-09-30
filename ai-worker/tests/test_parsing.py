from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from uuid import uuid4
import json

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject
import pytest
from fastapi.testclient import TestClient

from researchhub_worker.contracts import SourceIngestCommand
from researchhub_worker.parsing import SourceParser, DownloadFailure, download
from researchhub_worker.parsing.common import ParserLimits, ParseFailure, check_office_archive
from researchhub_worker.processor import IdempotentSourceIngestProcessor
from researchhub_worker.server import create_app


def command(kind='PDF'):
    data = json.loads((Path(__file__).resolve().parents[2] / 'contracts/processing/v4/source-ingest-request.json').read_text())
    return SourceIngestCommand.model_validate(data | {'sourceType': kind, 'jobId': str(uuid4())})


def pdf_bytes(texts, image_pages=(), encrypted=False):
    writer = PdfWriter()
    for index, text in enumerate(texts):
        page = writer.add_blank_page(width=600, height=800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                                 NameObject('/BaseFont'): NameObject('/Helvetica')})
        resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        if index in image_pages:
            image = DecodedStreamObject()
            image.set_data(b'\x00\x00\x00')
            image.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                NameObject('/Width'): NumberObject(1), NameObject('/Height'): NumberObject(1),
                NameObject('/BitsPerComponent'): NumberObject(8), NameObject('/ColorSpace'): NameObject('/DeviceRGB')})
            resources[NameObject('/XObject')] = DictionaryObject({NameObject('/I1'): writer._add_object(image)})
        page[NameObject('/Resources')] = resources
        content = DecodedStreamObject()
        content.set_data(f'BT /F1 12 Tf 10 700 Td ({text}) Tj ET'.encode('ascii') if text else b'')
        page[NameObject('/Contents')] = writer._add_object(content)
    if encrypted:
        writer.encrypt('secret')
    output = BytesIO(); writer.write(output)
    return output.getvalue()


def docx_bytes(empty=False):
    document = Document()
    if not empty:
        document.add_paragraph('Before')
        document.add_heading('Theory', 1)
        document.add_paragraph('Paragraph with Unicode: λ 😀')
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [('Name', 'Value'), ('mass', '3')]):
            for cell, value in zip(row.cells, values): cell.text = value
        document.add_heading('Details', 2)
        document.add_paragraph('After table')
        document.add_heading('Conclusion', 1)
    output = BytesIO(); document.save(output)
    return output.getvalue()


def xlsx_bytes(rows=5):
    workbook = Workbook()
    sheet = workbook.active; sheet.title = 'Measurements'
    sheet.append(['name', 'value', 'formula', 'flag', 'date'])
    for index in range(rows): sheet.append([f'row-{index}', index, '=B2*2', True, datetime(2026, 9, 29)])
    hidden = workbook.create_sheet('Hidden'); hidden.sheet_state = 'hidden'; hidden.append(['secret', '#DIV/0!'])
    very = workbook.create_sheet('Internal'); very.sheet_state = 'veryHidden'; very.append([None, 'x'])
    workbook.create_sheet('Empty')
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


def parse(kind, data, limits=None):
    return SourceParser(limits=limits, downloader=lambda *_: data)(command(kind))


def test_lecture_pdf_pages_provenance_hash_and_determinism():
    work = command(); data = pdf_bytes(['Lecture 1', 'Lecture 2', ''])
    parser = SourceParser(downloader=lambda *_: data)
    first = parser(work); second = parser(work)
    assert first == second
    assert first.status == 'SUCCEEDED'
    assert [unit.text for unit in first.chunks] == ['Lecture 1', 'Lecture 2', '']
    assert [unit.page_number for unit in first.chunks] == [1, 2, 3]
    assert all(unit.source_id == work.source_id and unit.parser_version == first.parser_version for unit in first.chunks)
    assert first.extraction_metadata.page_count == 3
    assert first.extraction_metadata.character_count == 18
    assert first.structure.pages[1].character_start == 9
    assert len(first.extraction_metadata.content_sha256) == 64


@pytest.mark.parametrize('kind,data,code', [
    ('PDF', b'%PDF-1.7\ncorrupted', 'DOCUMENT_PARSE_FAILED'),
    ('PDF', pdf_bytes([''], image_pages=[0]), 'OCR_REQUIRED'),
    ('PDF', pdf_bytes(['']), 'EMPTY_DOCUMENT'),
    ('PDF', pdf_bytes(['text'], encrypted=True), 'PDF_ENCRYPTED'),
    ('DOCX', b'PKbad', 'DOCUMENT_PARSE_FAILED'),
    ('DOCX', docx_bytes(empty=True), 'EMPTY_DOCUMENT'),
    ('XLSX', b'PKbad', 'DOCUMENT_PARSE_FAILED'),
    ('TXT', b'\xff', 'DOCUMENT_PARSE_FAILED'),
    ('TXT', b' ', 'EMPTY_DOCUMENT'),
])
def test_safe_failures(kind, data, code):
    result = parse(kind, data)
    assert result.status == 'FAILED' and result.failure.code == code
    assert result.chunks == [] and result.extraction_metadata is None


def test_mixed_pdf_preserves_blank_page_and_warns():
    result = parse('PDF', pdf_bytes(['Text', ''], image_pages=[1]))
    assert result.status == 'SUCCEEDED'
    assert result.chunks[1].page_number == 2
    assert result.warnings[0].startswith('OCR_REQUIRED')


def test_docx_body_order_sections_and_table_locations_without_office():
    work = command('DOCX'); parser = SourceParser(downloader=lambda *_: docx_bytes())
    result = parser(work)
    assert result == parser(work)
    assert result.status == 'SUCCEEDED'
    assert [chunk.location.kind for chunk in result.chunks] == ['PARAGRAPH', 'HEADING', 'PARAGRAPH', 'TABLE', 'HEADING', 'PARAGRAPH', 'HEADING']
    assert result.chunks[3].text == 'Name\tValue\nmass\t3'
    assert result.chunks[0].page_number is None
    assert result.chunks[5].section_id == 'section-4'
    assert result.structure.sections[1].parent_section_id == 'section-1'
    assert result.structure.sections[0].character_end == result.chunks[6].character_start
    assert result.extraction_metadata.character_count == sum(len(chunk.text) for chunk in result.chunks)


def test_multisheet_metadata_hidden_types_formulas_and_bounded_sampling():
    work = command('XLSX'); data = xlsx_bytes()
    parser = SourceParser(downloader=lambda *_: data)
    result = parser(work)
    assert result == parser(work)
    assert result.status == 'SUCCEEDED'
    sheets = result.workbook.sheets
    assert [sheet.name for sheet in sheets] == ['Measurements', 'Hidden', 'Internal', 'Empty']
    assert [sheet.state for sheet in sheets] == ['visible', 'hidden', 'veryHidden', 'visible']
    assert sheets[0].used_range == 'A1:E6' and sheets[0].row_count_estimate == 6
    assert sheets[0].header_candidate == ['name', 'value', 'formula', 'flag', 'date']
    assert sheets[0].formula_presence is True and sheets[0].formula_scan_complete is True
    assert sheets[0].columns[2].values[1] == '=B2*2'
    assert sheets[0].columns[3].data_types == ['boolean', 'text']
    assert sheets[0].columns[4].data_types == ['date', 'text']
    assert sheets[1].columns[1].data_types == ['error']
    assert sheets[3].header_row is None
    limited = parse('XLSX', data, ParserLimits(xlsx_rows=2, xlsx_columns=2, xlsx_samples=1))
    assert limited.workbook.sheets[0].sampled_rows == 2
    assert limited.workbook.sheets[0].truncated is True
    assert limited.workbook.sheets[0].formula_presence is None
    assert limited.workbook.sheets[0].formula_scan_complete is False
    assert all(len(column.values) <= 1 for column in limited.workbook.sheets[0].columns)
    assert limited.warnings


def test_archive_refuses_disguised_macros_and_zip_expansion():
    data = xlsx_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, 'w') as target:
        for name in source.namelist(): target.writestr(name, source.read(name))
        target.writestr('xl/vbaProject.bin', b'macro')
    assert parse('XLSX', output.getvalue()).failure.code == 'UNSUPPORTED_MACROS'
    with pytest.raises(ParseFailure, match='archive'):
        check_office_archive(data, ParserLimits(max_zip_bytes=1))


@pytest.mark.parametrize('kind,data,limits', [
    ('PDF', pdf_bytes(['one', 'two']), ParserLimits(max_pages=1)),
    ('DOCX', docx_bytes(), ParserLimits(max_characters=5)),
    ('DOCX', docx_bytes(), ParserLimits(max_units=1)),
    ('XLSX', xlsx_bytes(), ParserLimits(max_sheets=1)),
])
def test_configured_limits_fail_safely(kind, data, limits):
    assert parse(kind, data, limits).failure.code == 'EXTRACTION_LIMIT_EXCEEDED'


def test_utf8_text_keeps_identity():
    result = parse('TXT', 'λ,😀\nvalue,2'.encode('utf-8-sig'))
    assert result.status == 'SUCCEEDED' and result.chunks[0].text == 'λ,😀\nvalue,2'


def test_runtime_limits_are_validated(monkeypatch):
    monkeypatch.setenv('AI_WORKER_XLSX_ROWS', '2')
    assert ParserLimits.from_env().xlsx_rows == 2
    monkeypatch.setenv('AI_WORKER_XLSX_ROWS', '0')
    with pytest.raises(ValueError): ParserLimits.from_env()
    monkeypatch.setenv('AI_WORKER_XLSX_ROWS', 'wrong')
    with pytest.raises(ValueError): ParserLimits.from_env()


def test_transient_download_is_not_cached():
    calls = 0
    def fail_then_download(*_):
        nonlocal calls
        calls += 1
        if calls == 1: raise DownloadFailure('Temporary access failed.')
        return pdf_bytes(['Recovered'])
    processor = IdempotentSourceIngestProcessor(SourceParser(downloader=fail_then_download))
    work = command()
    with pytest.raises(DownloadFailure): processor.process(work)
    assert processor.process(work).result.status == 'SUCCEEDED'
    assert processor.process(work).duplicate is True and calls == 2


def test_http_ingestion_returns_real_extraction_and_safe_malformed_failure():
    token = 'test-token-with-at-least-32-characters'
    client = TestClient(create_app(IdempotentSourceIngestProcessor(SourceParser(downloader=lambda *_: docx_bytes())), token))
    response = client.post('/internal/jobs/source-ingest', json=command('DOCX').model_dump(mode='json', by_alias=True), headers={'Authorization': 'Bearer '+token})
    assert response.status_code == 200 and response.json()['chunks'][3]['location']['kind'] == 'TABLE'
    client = TestClient(create_app(IdempotentSourceIngestProcessor(SourceParser(downloader=lambda *_: b'%PDF broken')), token))
    response = client.post('/internal/jobs/source-ingest', json=command().model_dump(mode='json', by_alias=True), headers={'Authorization': 'Bearer '+token})
    assert response.status_code == 200 and response.json()['failure']['code'] == 'DOCUMENT_PARSE_FAILED'


def test_download_bounds_expiry_transport_and_redirects():
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    from threading import Thread
    from pydantic import AnyHttpUrl
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == '/redirect':
                self.send_response(302); self.send_header('Location', '/file'); self.end_headers(); return
            self.send_response(200); self.end_headers(); self.wfile.write(b'hello')
        def log_message(self, *_): pass
    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = Thread(target=server.serve_forever, daemon=True); thread.start()
    try:
        work = command('TXT')
        access = work.file_access.model_copy(update={'url': AnyHttpUrl(f'http://127.0.0.1:{server.server_port}/file')})
        work = work.model_copy(update={'file_access': access})
        assert download(work, ParserLimits()) == b'hello'
        with pytest.raises(ParseFailure): download(work, ParserLimits(max_file_bytes=1))
        expired = work.model_copy(update={'file_access': access.model_copy(update={'expires_at': datetime(2000, 1, 1, tzinfo=timezone.utc)})})
        with pytest.raises(DownloadFailure, match='expired'): download(expired, ParserLimits())
        redirect = work.model_copy(update={'file_access': access.model_copy(update={'url': AnyHttpUrl(f'http://127.0.0.1:{server.server_port}/redirect')})})
        with pytest.raises(DownloadFailure) as failure: download(redirect, ParserLimits())
        assert '127.0.0.1' not in str(failure.value)
    finally:
        server.shutdown(); server.server_close(); thread.join()
    with pytest.raises(DownloadFailure): download(work, ParserLimits())


def test_bounded_idempotency_cache_allows_reexecution_after_eviction():
    from researchhub_worker.contracts import SourceIngestResult
    processor = IdempotentSourceIngestProcessor(SourceIngestResult.empty_success)
    first = command(); processor.process(first)
    for _ in range(16): processor.process(command())
    assert len(processor._completed) == 16
    assert processor.process(first).duplicate is False


@pytest.mark.parametrize('origin,target,url,expected', [
    (None, None, 'https://cloud.example/file?sig=x', 'https://cloud.example/file?sig=x'),
    ('http://127.0.0.1:10000', 'http://azurite:10000', 'http://127.0.0.1:10000/devstoreaccount1/blob?sig=a%2Bb', 'http://azurite:10000/devstoreaccount1/blob?sig=a%2Bb'),
    ('http://127.0.0.1:10000', 'http://azurite:10000', 'https://cloud.example/file?sig=x', 'https://cloud.example/file?sig=x'),
])
def test_local_origin_translation_preserves_signed_path_query(monkeypatch, origin, target, url, expected):
    from researchhub_worker.parsing import _download_url
    for name, value in [('AI_WORKER_BLOB_URL_ORIGIN', origin), ('AI_WORKER_BLOB_URL_TARGET_ORIGIN', target)]:
        if value is None: monkeypatch.delenv(name, raising=False)
        else: monkeypatch.setenv(name, value)
    assert _download_url(url) == expected


@pytest.mark.parametrize('origin,target', [('http://host', None), ('file:///tmp', 'http://azurite'), ('http://host/path', 'http://azurite'), ('http://user:pass@host', 'http://azurite')])
def test_invalid_origin_translation_fails_safely(monkeypatch, origin, target):
    from researchhub_worker.parsing import _download_url
    monkeypatch.setenv('AI_WORKER_BLOB_URL_ORIGIN', origin)
    if target: monkeypatch.setenv('AI_WORKER_BLOB_URL_TARGET_ORIGIN', target)
    else: monkeypatch.delenv('AI_WORKER_BLOB_URL_TARGET_ORIGIN', raising=False)
    with pytest.raises(DownloadFailure): _download_url('http://host/file')


def test_pdf_image_detection_handles_indirect_resources_and_nested_forms():
    from researchhub_worker.parsing.pdf import _has_image
    writer = PdfWriter()
    image = DictionaryObject({NameObject('/Subtype'): NameObject('/Image')})
    inner = DictionaryObject({NameObject('/XObject'): writer._add_object(DictionaryObject({NameObject('/Image'): writer._add_object(image)}))})
    form = DictionaryObject({NameObject('/Subtype'): NameObject('/Form'), NameObject('/Resources'): writer._add_object(inner)})
    resources = writer._add_object(DictionaryObject({NameObject('/XObject'): writer._add_object(DictionaryObject({NameObject('/Form'): writer._add_object(form)}))}))
    assert _has_image(resources)


def test_docx_nested_table_keeps_cell_order_and_ignores_empty_blocks():
    document = Document()
    document.add_paragraph('')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Before nested'
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'Nested value'
    output = BytesIO(); document.save(output)
    result = parse('DOCX', output.getvalue())
    assert result.status == 'SUCCEEDED'
    assert result.chunks[0].text == 'Before nested\nNested value'
    assert result.chunks[0].location.block_index == 1


def test_csv_preserves_quoted_cells_unicode_newlines_and_preview_limits():
    result = parse('CSV', 'name;note;formula\nλ;"line 1\nline 2";=2+2\nlast;"has;separator";😀'.encode('utf-8-sig'),
                   ParserLimits(preview_rows=2, xlsx_rows=2, xlsx_columns=3))
    assert result.status == 'SUCCEEDED'
    sheet = result.workbook.sheets[0]
    assert sheet.header_candidate == ['name', 'note', 'formula']
    assert sheet.preview_rows[1].cells == ['λ', 'line 1\nline 2', '=2+2']
    assert len(sheet.preview_rows) == 2 and sheet.truncated
    assert sheet.row_count_estimate is None and result.warnings
    assert 'text' in sheet.columns[2].data_types
    assert parse('CSV', b'one').workbook.sheets[0].preview_rows[0].cells == ['one']
    assert parse('CSV', b' ').failure.code == 'EMPTY_DOCUMENT'
    assert parse('CSV', b'a,b\n"unterminated').failure.code == 'DOCUMENT_PARSE_FAILED'


def test_xlsx_preview_is_bounded_keeps_formulas_as_data_and_row_numbers():
    result = parse('XLSX', xlsx_bytes(10), ParserLimits(preview_rows=2, xlsx_rows=4))
    sheet = result.workbook.sheets[0]
    assert [row.row_number for row in sheet.preview_rows] == [1, 2]
    assert sheet.preview_rows[1].cells[2] == '=B2*2'
    assert len(sheet.preview_rows[1].cells) == 5
    assert sheet.sampled_rows == 4 and sheet.truncated
    assert result.workbook.preview_row_limit == 2


def test_csv_header_candidate_skips_empty_records_without_reordering_preview():
    result = parse('CSV', b',\nname,value\nAda,3\n')
    assert result.status == 'SUCCEEDED'
    sheet = result.workbook.sheets[0]
    assert sheet.header_row == 2 and sheet.header_candidate == ['name', 'value']
    assert sheet.preview_rows[0].row_number == 1 and sheet.preview_rows[0].cells == ['', '']
    assert sheet.preview_rows[2].cells == ['Ada', '3']
