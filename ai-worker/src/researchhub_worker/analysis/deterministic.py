"""Narrow offline fixture: generate an impedance program, never derive results from sampled rows."""
import re
import unicodedata

ALIASES = {
    'frequency': {'f', 'frequency', 'czestotliwosc'},
    'voltage': {'u', 'voltage', 'napiecie'},
    'current': {'i', 'current', 'prad'},
}
UNITS = {'frequency': {'hz': 1, 'khz': 1000, 'mhz': 1_000_000},
         'voltage': {'v': 1, 'mv': 0.001}, 'current': {'a': 1, 'ma': 0.001, 'ua': 0.000001}}


def column_role(name):
    normal = unicodedata.normalize('NFKD', name.replace('µ', 'u').replace('μ', 'u'))
    normal = ''.join(character for character in normal if not unicodedata.combining(character)).lower()
    tokens = [token for token in re.split(r'[^a-z]+', normal) if token]
    for role, aliases in ALIASES.items():
        if tokens and tokens[0] in aliases:
            if len(tokens) == 1:
                return role, 1, True
            if len(tokens) == 2 and tokens[1] in UNITS[role]:
                return role, UNITS[role][tokens[1]], False
    return None


def impedance_plan(request):
    if 'impedan' not in request.request.instruction.lower() or len(request.inputs) != 1:
        return None
    inspected = request.inputs[0]
    selection = inspected.selection
    candidates = []
    for sheet in inspected.preview.sheets:
        if selection.sheet_name is not None and selection.sheet_name != sheet.name or sheet.header_row is None:
            continue
        roles = {}
        for column in sheet.columns:
            if selection.columns is not None and column.index not in selection.columns:
                continue
            found = column_role(column.name)
            if found:
                roles.setdefault(found[0], []).append((column.index, found[1], found[2]))
        if set(roles) == set(ALIASES) and all(len(values) == 1 for values in roles.values()):
            candidates.append((sheet, {role: values[0] for role, values in roles.items()}))
    if len(candidates) != 1:
        return None
    sheet, roles = candidates[0]
    version = str(selection.source_version_id)
    columns = [roles[role][0] for role in ('frequency', 'voltage', 'current')]
    assumptions = ['Frequency must be positive; current must be nonzero. Invalid, missing and formula cells stop the run.',
                   'The table contains signed Z = U / I; the chart plots its magnitude |Z|.']
    bare = [role for role, (_, _, assumed) in roles.items() if assumed]
    if bare:
        assumptions.append('Bare column labels use SI units: frequency Hz, voltage V and current A; verify before execution.')
    warnings = ['Deterministic offline impedance fixture; numeric results are produced only by isolated execution.']
    if bare:
        warnings.append('Column units were assumed from bare labels; choose explicitly labelled units when available.')
    path = '/inputs/' + version + '.' + inspected.preview.format.lower()
    code = _code(path, inspected.preview.format, sheet.name, sheet.header_row,
                 [roles[role][0] - 1 for role in ('frequency', 'voltage', 'current')],
                 [roles[role][1] for role in ('frequency', 'voltage', 'current')])
    return {'schemaVersion': '1.0', 'summary': 'Calculate impedance Z = U / I versus frequency from the selected immutable dataset.',
            'inputs': [{'sourceVersionId': version, 'sheetName': sheet.name, 'requiredColumns': columns}],
            'transformations': [{'name': 'impedance', 'description': 'Convert explicitly labelled units to SI, then calculate Z = U / I for each row.',
                                 'sourceVersionId': version, 'sheetName': sheet.name, 'columns': columns}],
            'statisticalOperations': [],
            'outputs': [{'kind': 'TABLE', 'name': 'impedance-table', 'description': 'Frequency, voltage, current and calculated impedance in SI units.', 'sourceVersionIds': [version]},
                        {'kind': 'CHART', 'name': 'impedance-chart', 'description': 'Impedance magnitude versus frequency with labelled axes.', 'sourceVersionIds': [version]}],
            'assumptions': assumptions, 'warnings': warnings, 'code': {'language': 'PYTHON', 'source': code}}


def _code(path, file_format, sheet, header_row, positions, factors):
    # repr emits inert Python literals. No preview value is interpolated into executable statements.
    return f'''# ResearchHub deterministic impedance fixture v1; run only inside the isolated executor.
import csv
import io
import json
import math
from pathlib import Path
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook

INPUT = Path({path!r})
FORMAT = {file_format!r}
SHEET = {sheet!r}
HEADER_ROW = {header_row!r}
POSITIONS = {positions!r}
FACTORS = {factors!r}

if FORMAT == 'CSV':
    text = INPUT.read_bytes().decode('utf-8-sig')
    try:
        delimiter = csv.Sniffer().sniff(text[:65536], delimiters=',;\\t|').delimiter
    except csv.Error:
        raise ValueError('CSV delimiter must be unambiguous') from None
    records = list(csv.reader(io.StringIO(text, newline=''), delimiter=delimiter, strict=True))
else:
    workbook = load_workbook(INPUT, read_only=True, data_only=False, keep_links=False)
    records = list(workbook[SHEET].iter_rows(values_only=True))
    workbook.close()
header = records[HEADER_ROW - 1]
rows = []
for number, record in enumerate(records[HEADER_ROW:], HEADER_ROW + 1):
    if not record or all(value is None or str(value).strip() == '' for value in record):
        continue
    if FORMAT == 'CSV' and len(record) != len(header):
        raise ValueError('Inconsistent CSV row width')
    if len(rows) >= 10000:
        raise ValueError('Result row limit exceeded')
    values = []
    for position, factor in zip(POSITIONS, FACTORS):
        value = record[position] if position < len(record) else None
        if value is None or isinstance(value, bool) or str(value).strip().startswith('='):
            raise ValueError('A required numeric cell is missing or contains a formula')
        try:
            converted = float(value) * factor
        except (ValueError, TypeError):
            raise ValueError('A required cell contains nonnumeric data') from None
        if not math.isfinite(converted):
            raise ValueError('Numeric values must be finite')
        values.append(converted)
    frequency, voltage, current = values
    if frequency <= 0 or current == 0:
        raise ValueError('Frequency must be positive and current nonzero')
    impedance = voltage / current
    if not math.isfinite(impedance):
        raise ValueError('Calculated impedance must be finite')
    rows.append(values + [impedance])
if not rows:
    raise ValueError('No measurement rows are available')
columns = ['Frequency (Hz)', 'Voltage (V)', 'Current (A)', 'Impedance (Ω)']
frame = pd.DataFrame(rows, columns=columns).sort_values(columns[0], kind='stable')
figure, axis = plt.subplots(figsize=(8, 5), layout='constrained')
axis.plot(frame.iloc[:, 0], np.abs(frame.iloc[:, 3]), marker='o', color='#63a4d9')
axis.set_xscale('log')
if (np.abs(frame.iloc[:, 3]) > 0).all():
    axis.set_yscale('log')
axis.set_xlabel('Frequency f (Hz)')
axis.set_ylabel('Impedance magnitude |Z| (Ω)')
axis.set_title('Impedance magnitude versus frequency')
axis.grid(True, which='both', alpha=0.25)
figure.savefig('/outputs/impedance.png', dpi=120, metadata={{'Software': 'ResearchHub sandbox 1.0.0'}})
plt.close(figure)
result = {{'schemaVersion': '1.0', 'outputs': [
    {{'name': 'impedance-table', 'kind': 'TABLE', 'columns': columns, 'rows': frame.values.tolist()}},
    {{'name': 'impedance-chart', 'kind': 'CHART', 'file': 'impedance.png'}}]}}
Path('/outputs/result.json').write_text(json.dumps(result, allow_nan=False, separators=(',', ':')), encoding='utf-8')
'''
