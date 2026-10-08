"""Project RC demo only: build an explicit program; never compute from preview samples."""


def rc_plan(request):
    instruction = request.request.instruction.lower()
    if 'rc' not in instruction or not ('tau' in instruction or 'time constant' in instruction) or len(request.inputs) != 1:
        return None
    inspected = request.inputs[0]
    if inspected.preview.format != 'XLSX':
        return None
    selection = inspected.selection
    candidates = []
    for sheet in inspected.preview.sheets:
        if sheet.header_row is None or selection.sheet_name is not None and sheet.name != selection.sheet_name:
            continue
        columns = {name: [c.index for c in sheet.columns if c.name == name and
                   (selection.columns is None or c.index in selection.columns)] for name in ('time_s', 'voltage_v')}
        if all(len(indices) == 1 for indices in columns.values()):
            candidates.append((sheet, [columns[name][0] for name in ('time_s', 'voltage_v')]))
    if len(candidates) != 1:
        return None
    sheet, columns = candidates[0]
    version = str(selection.source_version_id)
    source = _code('/inputs/' + version + '.xlsx', sheet.name, sheet.header_row, [c - 1 for c in columns])
    step = {'name': 'log-linear-fit', 'description': 'Least-squares log(V) versus time from all positive finite immutable measurements; tau = -1/slope.',
            'sourceVersionId': version, 'sheetName': sheet.name, 'columns': columns}
    return {'schemaVersion': '1.0', 'summary': 'Estimate RC discharge tau and plot measured versus fitted voltage.',
            'inputs': [{'sourceVersionId': version, 'sheetName': sheet.name, 'requiredColumns': columns}],
            'transformations': [], 'statisticalOperations': [step],
            'outputs': [{'kind': kind, 'name': name, 'description': description, 'sourceVersionIds': [version]}
                        for kind, name, description in [
                            ('TABLE', 'rc-fit-table', 'Time, measured voltage, fitted voltage and residual from all rows.'),
                            ('CHART', 'rc-discharge-chart', 'Measured and fitted discharge curves in seconds and volts.'),
                            ('TEXT', 'rc-fit-summary', 'Estimated tau and log-space fit quality with assumptions.')]],
            'assumptions': ['Zero voltage offset; log-linear least squares pools the repeated trials.',
                            'Additive voltage noise can bias a log fit; compare residuals and component tolerances.'],
            'warnings': ['Deterministic project-authored RC fixture. Numeric results exist only after real isolated execution.'],
            'code': {'language': 'PYTHON', 'source': source}}


def _code(path, sheet, header, positions):
    return f'''# ResearchHub synthetic RC fixture v1: execute only in the isolated sandbox.
import json
import math
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from openpyxl import load_workbook

workbook = load_workbook({path!r}, read_only=True, data_only=False, keep_links=False)
records = list(workbook[{sheet!r}].iter_rows(values_only=True))
workbook.close()
rows = []
for record in records[{header!r}:]:
    if all(value is None for value in record):
        continue
    if len(rows) >= 10000:
        raise ValueError('Measurement row limit exceeded')
    values = []
    for index in {positions!r}:
        value = record[index] if index < len(record) else None
        if value is None or isinstance(value, bool) or str(value).startswith('='):
            raise ValueError('Missing or formula measurement')
        converted = float(value)
        if not math.isfinite(converted):
            raise ValueError('Measurements must be finite')
        values.append(converted)
    if values[0] < 0 or values[1] <= 0:
        raise ValueError('Time must be nonnegative and voltage positive')
    rows.append(values)
if len(rows) < 3:
    raise ValueError('At least three measurements are required')
data = np.asarray(rows, dtype=float)
time, voltage = data[:, 0], data[:, 1]
if np.unique(time).size < 3:
    raise ValueError('At least three distinct times are required')
log_voltage = np.log(voltage)
slope, intercept = np.polyfit(time, log_voltage, 1)
if slope >= 0:
    raise ValueError('RC discharge needs a negative fitted slope')
tau = float(-1 / slope)
fitted = np.exp(intercept + slope * time)
variance = float(np.sum((log_voltage - np.mean(log_voltage)) ** 2))
if variance == 0:
    raise ValueError('Constant voltage does not define a discharge fit')
r_squared = float(1 - np.sum((log_voltage - (intercept + slope * time)) ** 2) / variance)
table = np.column_stack((time, voltage, fitted, voltage - fitted))
columns = ['Time (s)', 'Measured voltage (V)', 'Fitted voltage (V)', 'Residual (V)']
order = np.argsort(time, kind='stable')
figure, axis = plt.subplots(figsize=(8, 5), layout='constrained')
axis.scatter(time, voltage, s=10, label='Measured', color='#16324f')
axis.plot(time[order], fitted[order], label='Fitted', color='#e07a24')
axis.set(xlabel='Time (s)', ylabel='Voltage (V)', title='RC discharge: measured versus fitted')
axis.legend()
axis.grid(alpha=0.25)
figure.savefig('/outputs/rc-discharge.png', dpi=120)
plt.close(figure)
summary = 'Estimated tau = %.6f s; log-space R squared = %.6f. Zero-offset log-linear fit; pooled repeated trials; additive noise may bias the estimate.' % (tau, r_squared)
result = {{'schemaVersion': '2.0', 'outputs': [
    {{'name': 'rc-fit-table', 'kind': 'TABLE', 'columns': columns, 'rows': table.tolist()}},
    {{'name': 'rc-fit-summary', 'kind': 'TEXT', 'text': summary}},
    {{'name': 'rc-discharge-chart', 'kind': 'CHART', 'file': 'rc-discharge.png',
      'title': 'RC discharge: measured versus fitted',
      'xAxis': {{'label': 'Time', 'unit': 's', 'scale': 'LINEAR'}},
      'yAxis': {{'label': 'Voltage', 'unit': 'V', 'scale': 'LINEAR'}},
      'series': [{{'name': name, 'tableName': 'rc-fit-table', 'xColumn': columns[0], 'yColumn': columns[index], 'yTransform': 'IDENTITY'}}
                 for name, index in [('Measured', 1), ('Fitted', 2)]]}}]}}
Path('/outputs/result.json').write_text(json.dumps(result, allow_nan=False, separators=(',', ':')), encoding='utf-8')
'''
