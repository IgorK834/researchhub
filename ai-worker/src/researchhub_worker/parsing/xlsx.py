"""Read-only workbook profiling. Formula strings are data and never evaluated."""
from datetime import date, datetime, time
from itertools import islice
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter
from ..contracts import ColumnSample, PreviewRow, SheetMetadata, UnitLocation, WorkbookMetadata
from .common import ParseFailure, TextBuilder, check_office_archive

PARSER_VERSION = 'openpyxl-3.1.5/rh-2'


def _value(cell):
    value = cell.value
    if value is None:
        return '', 'empty'
    if cell.data_type == 'f':
        return str(value)[:500], 'formula'
    if cell.data_type == 'e':
        return str(value)[:500], 'error'
    if isinstance(value, bool):
        return str(value).lower(), 'boolean'
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(), 'date'
    if isinstance(value, (int, float)):
        return str(value), 'number'
    return str(value)[:500], 'text'


def parse_xlsx(data, source_id, limits):
    check_office_archive(data, limits)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_vba=False, keep_links=False)
    builder = TextBuilder(source_id, PARSER_VERSION, limits)
    sheets, warnings = [], []
    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise ParseFailure('EXTRACTION_LIMIT_EXCEEDED', 'The workbook exceeds the configured sheet limit.')
        for sheet in workbook.worksheets:
            # Dimensions are producer-supplied estimates, not an exhaustive scan. Reset before
            # bounded iteration so incorrect dimensions cannot hide sampled cells.
            rows, columns = sheet.max_row, sheet.max_column
            used_range = sheet.calculate_dimension() if rows is not None and columns is not None else None
            sheet.reset_dimensions()
            sampled = list(islice(sheet.iter_rows(max_col=limits.xlsx_columns), limits.xlsx_rows + 1))
            extra_row = len(sampled) > limits.xlsx_rows
            sampled = sampled[:limits.xlsx_rows]
            found_formula = False
            header, header_row = [], None
            samples = [[] for _ in range(limits.xlsx_columns)]
            types = [set() for _ in range(limits.xlsx_columns)]
            lines, preview = [], []
            for row_number, row in enumerate(sampled, 1):
                values = []
                for index, cell in enumerate(row):
                    value, kind = _value(cell)
                    values.append(value)
                    if kind == 'formula':
                        found_formula = True
                    if kind != 'empty':
                        types[index].add(kind)
                        if len(samples[index]) < limits.xlsx_samples:
                            samples[index].append(value)
                while values and not values[-1]:
                    values.pop()
                if row_number <= limits.preview_rows:
                    preview.append(PreviewRow(row_number=row_number, cells=values.copy()))
                if values:
                    if header_row is None:
                        header, header_row = values, row_number
                    lines.append(f'{row_number}: ' + '\t'.join(values))
            column_samples = [ColumnSample(column_number=index+1, values=values, data_types=sorted(types[index]))
                              for index, values in enumerate(samples) if values]
            truncated = extra_row or (rows is not None and rows > limits.xlsx_rows) or (columns is None or columns > limits.xlsx_columns)
            complete = not truncated
            sheets.append(SheetMetadata(preview_rows=preview, name=sheet.title, state=sheet.sheet_state, used_range=used_range,
                row_count_estimate=rows, column_count=columns, header_candidate=header, header_row=header_row,
                columns=column_samples, sampled_rows=len(sampled), truncated=truncated,
                formula_presence=True if found_formula else (False if complete else None), formula_scan_complete=complete))
            sampled_range = f'A1:{get_column_letter(limits.xlsx_columns)}{max(1, len(sampled))}'
            builder.add('\n'.join(lines), UnitLocation(kind='SHEET', sheet_name=sheet.title, cell_range=sampled_range))
        if any(sheet.truncated for sheet in sheets):
            warnings.append('WORKBOOK_SAMPLED: Row/column limits apply; dimensions are estimates and formula detection may be incomplete.')
    finally:
        workbook.close()
    metadata = WorkbookMetadata(preview_row_limit=limits.preview_rows, sheets=sheets, row_limit=limits.xlsx_rows, column_limit=limits.xlsx_columns,
                                sample_limit=limits.xlsx_samples)
    return builder, [], [], metadata, warnings, None, None
